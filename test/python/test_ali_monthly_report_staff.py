from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import ModuleType
from typing import Any, Callable
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
MODULAR_TOOL_DIR = ROOT / "tools" / "ali-monthly-report"
LEGACY_TOOL_DIR = ROOT / "tools" / "ali-monthly-report-email"


def load_legacy_app() -> ModuleType:
    module_name = "nantong_ali_monthly_report_legacy_app_test"
    existing = sys.modules.get(module_name)
    if existing is not None:
        return existing

    sys.path.insert(0, str(LEGACY_TOOL_DIR))
    try:
        spec = importlib.util.spec_from_file_location(
            module_name,
            LEGACY_TOOL_DIR / "app.py",
        )
        if spec is None or spec.loader is None:
            raise RuntimeError("无法加载阿里月报旧程序")
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.remove(str(LEGACY_TOOL_DIR))


sys.path.insert(0, str(MODULAR_TOOL_DIR))
try:
    from ali_monthly_reports import staff as modular_staff
finally:
    sys.path.remove(str(MODULAR_TOOL_DIR))


Writer = Callable[[list[dict[str, Any]], list[dict[str, Any]], str, Path], dict[str, int]]


class StaffRosterWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.legacy_app = load_legacy_app()

    def test_both_generators_keep_formal_roster_and_group_all_monthly_changes(self) -> None:
        active_records = [
            {
                "姓名": "正式在职",
                "岗位": "暖通工程师",
                "岗位性质": "正式岗",
                "入职日期": datetime(2025, 4, 1),
            },
            {
                "姓名": "备岗在职",
                "岗位": "电气备岗",
                "岗位性质": "备岗",
                "入职日期": datetime(2025, 5, 1),
            },
        ]
        all_records = [
            {
                "姓名": "实习新人",
                "岗位": "运维实习生",
                "岗位性质": "实习生",
                "入职日期": datetime(2026, 5, 3),
            },
            {
                "姓名": "备岗新人",
                "岗位": "电气备岗",
                "岗位性质": "备岗",
                "入职日期": datetime(2026, 5, 5),
            },
            {
                "姓名": "正式新人",
                "岗位": "暖通工程师",
                "岗位性质": "正式岗",
                "入职日期": datetime(2026, 5, 7),
            },
            {
                "姓名": "实习离职",
                "岗位": "运维实习生",
                "岗位性质": "实习生",
                "离职/异动情况": True,
                "离职/异动日期": datetime(2026, 5, 9),
            },
            {
                "姓名": "正式离职",
                "岗位": "电气工程师",
                "岗位性质": "正式岗",
                "离职/异动情况": True,
                "离职/异动日期": datetime(2026, 5, 11),
            },
        ]
        writers: list[tuple[str, Writer]] = [
            ("modular", modular_staff.write_staff_roster_workbook),
            ("legacy", self.legacy_app.write_staff_roster_workbook),
        ]

        for label, writer in writers:
            with self.subTest(generator=label), tempfile.TemporaryDirectory(
                ignore_cleanup_errors=True,
            ) as temp_dir:
                output_path = Path(temp_dir) / f"{label}-staff.xlsx"
                counts = writer(active_records, all_records, "2026-05", output_path)

                self.assertEqual(counts["active_count"], 1)
                self.assertEqual(counts["new_hire_count"], 3)
                self.assertEqual(counts["leaver_count"], 2)

                workbook = load_workbook(output_path, read_only=False, data_only=True)
                try:
                    active_sheet = workbook["在职人员基础信息"]
                    self.assertEqual(active_sheet["A2"].value, "正式在职")
                    self.assertIsNone(active_sheet["A3"].value)

                    changes_sheet = workbook["本月人员变动"]
                    self.assertEqual(changes_sheet["A1"].value, "实习生及备岗")
                    self.assertIn("A1:D1", {str(item) for item in changes_sheet.merged_cells.ranges})
                    self.assertEqual(
                        [changes_sheet.cell(2, column).value for column in range(1, 5)],
                        ["本月新增人员姓名", "岗位", "离职人员姓名", "岗位"],
                    )
                    self.assertEqual(
                        [changes_sheet.cell(3, column).value for column in range(1, 5)],
                        ["实习新人", "运维实习生", "实习离职", "运维实习生"],
                    )
                    self.assertEqual(
                        [changes_sheet.cell(4, column).value for column in range(1, 5)],
                        ["备岗新人", "电气备岗", None, None],
                    )
                    self.assertIsNone(changes_sheet["A5"].value)
                    self.assertEqual(changes_sheet["A6"].value, "示例人员004员")
                    self.assertIn("A6:D6", {str(item) for item in changes_sheet.merged_cells.ranges})
                    self.assertEqual(
                        [changes_sheet.cell(7, column).value for column in range(1, 5)],
                        ["本月新增人员姓名", "岗位", "离职人员姓名", "岗位"],
                    )
                    self.assertEqual(
                        [changes_sheet.cell(8, column).value for column in range(1, 5)],
                        ["正式新人", "暖通工程师", "正式离职", "电气工程师"],
                    )
                finally:
                    workbook.close()


class LegacyLocalFileActionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.legacy_app = load_legacy_app()

    def test_local_file_routes_open_reveal_and_reject_unsafe_paths(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            output_dir = Path(temp_dir)
            report_path = output_dir / "monthly report.xlsx"
            unsafe_path = output_dir / "unsafe.cmd"
            report_path.write_bytes(b"report")
            unsafe_path.write_bytes(b"unsafe")

            client = self.legacy_app.app.test_client()
            with (
                patch.object(
                    self.legacy_app,
                    "load_config",
                    return_value={"download_folder": str(output_dir)},
                ),
                patch.object(
                    self.legacy_app.os,
                    "startfile",
                    create=True,
                ) as startfile,
                patch.object(self.legacy_app.subprocess, "Popen") as popen,
            ):
                open_response = client.post(
                    "/api/local_file/open",
                    json={"file_name": report_path.name},
                )
                reveal_response = client.post(
                    "/api/local_file/reveal",
                    json={"file_name": report_path.name},
                )
                traversal_response = client.post(
                    "/api/local_file/open",
                    json={"file_name": "../outside.xlsx"},
                )
                unsafe_response = client.post(
                    "/api/local_file/open",
                    json={"file_name": unsafe_path.name},
                )

            self.assertEqual(open_response.status_code, 200)
            self.assertIn("已提交系统打开请求", open_response.get_json()["message"])
            startfile.assert_called_once_with(str(report_path.resolve()))

            self.assertEqual(reveal_response.status_code, 200)
            self.assertIn("已提交资源管理器定位请求", reveal_response.get_json()["message"])
            popen.assert_called_once_with(
                ["explorer", f"/select,{report_path.resolve()}"],
            )

            self.assertEqual(traversal_response.status_code, 400)
            self.assertEqual(unsafe_response.status_code, 400)
            self.assertIn("仅支持打开或定位 XLSX、ZIP 文件", unsafe_response.get_json()["message"])

    def test_legacy_page_has_immediate_guarded_file_actions(self) -> None:
        template = (
            LEGACY_TOOL_DIR / "templates" / "index.html"
        ).read_text(encoding="utf-8")

        self.assertIn("file-item-actions", template)
        self.assertIn("打开文件", template)
        self.assertIn("定位文件", template)
        self.assertIn("generatedFileActionInFlight", template)
        self.assertIn("/api/local_file/${action}", template)
        self.assertLess(
            template.index("generatedFileActionInFlight.add(actionKey)"),
            template.index("await fetch(`/api/local_file/${action}`"),
        )


if __name__ == "__main__":
    unittest.main()
