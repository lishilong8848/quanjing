from __future__ import annotations

import calendar
import re
from collections import defaultdict
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl.cell.cell import MergedCell
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROW_LINE_SHEET = "机列超215功率统计"
LINE_HEAD_SHEET = "列头柜超107.5功率统计"
CABINET_SHEET = "机柜超18KW统计"
BRANCH_SHEET = "单支路超6.25KW功率"
INFO_SHEET = "说明"
ROW_LINE_APPENDIX_SHEET = "机列超功率附表"
LINE_HEAD_APPENDIX_SHEET = "列头柜超功率附表"
CABINET_APPENDIX_SHEET = "机柜超功率附表"

MAIN_SHEETS = [
    ROW_LINE_SHEET,
    LINE_HEAD_SHEET,
    CABINET_SHEET,
    BRANCH_SHEET,
]

HOUR_RE = re.compile(r"[-+]?\d+(?:\.\d+)?")
DEFAULT_FONT_NAME = "微软雅黑"
ROW_LINE_POWER_THRESHOLD = 215.0
E_BUILDING_ROW_LINE_POWER_THRESHOLD = 235.0
LINE_HEAD_POWER_THRESHOLD = 107.5
E_BUILDING_LINE_HEAD_POWER_THRESHOLD = 117.5


def generate_over_power_report(
    rows_by_kind: Mapping[str, Sequence[Mapping[str, Any]]],
    year: str,
    month_number: int,
    out: Path,
    data_center: str = "EA118",
    template_path: Path | None = None,
) -> dict[str, Any]:
    month_start, month_end = month_bounds(str(year), month_number)
    filtered = {
        kind: [dict(row) for row in rows if date_in_range(parse_date(row.get("数据时间")), month_start, month_end)]
        for kind, rows in rows_by_kind.items()
    }
    all_dates = sorted(
        parsed
        for rows in filtered.values()
        for row in rows
        if (parsed := parse_date(row.get("数据时间"))) is not None
    )
    if not all_dates:
        raise RuntimeError(f"没有找到 {year}年{month_number}月 的超功率源记录")

    date_label = f"{format_day(all_dates[0])}-{format_day(all_dates[-1])}"
    row_line_source_rows = filter_building_power_threshold(
        filtered.get("row_line", []),
        "功率",
        default_threshold=ROW_LINE_POWER_THRESHOLD,
        e_building_threshold=E_BUILDING_ROW_LINE_POWER_THRESHOLD,
    )
    line_head_source_rows = filter_building_power_threshold(
        filtered.get("line_head", []),
        "功率",
        default_threshold=LINE_HEAD_POWER_THRESHOLD,
        e_building_threshold=E_BUILDING_LINE_HEAD_POWER_THRESHOLD,
    )
    row_line_rows = aggregate_simple_rows(
        row_line_source_rows,
        key_fields=("楼栋", "房间", "机列"),
        power_field="功率",
        output_fields=("机房", "楼栋", "房间", "机列", "功率", "次数", "时长", "备注"),
        date_label=date_label,
        data_center=data_center,
    )
    line_head_rows = aggregate_simple_rows(
        line_head_source_rows,
        key_fields=("楼栋", "房间", "机列"),
        power_field="功率",
        output_fields=(
            "机房",
            "楼栋",
            "房间",
            "机列",
            "功率",
            "对侧机列",
            "对侧机列最大功率",
            "次数",
            "时长",
            "备注",
        ),
        date_label=date_label,
        data_center=data_center,
    )
    cabinet_groups = aggregate_cabinet_rows(filtered.get("cabinet", []), date_label=date_label, data_center=data_center)
    branch_rows = aggregate_branch_rows(filtered.get("branch", []), date_label=date_label, data_center=data_center)

    out.parent.mkdir(parents=True, exist_ok=True)
    if template_path is not None:
        wb = load_workbook(template_path)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    write_main_sheet(
        wb,
        ROW_LINE_SHEET,
        "EA118机列超功率告警统计表",
        ["序号", "数据时间", "机房", "楼栋", "房间", "机列", "功率", "次数", "时长", "备注"],
        rows_to_matrix(row_line_rows, ["序号", "数据时间", "机房", "楼栋", "房间", "机列", "功率", "次数", "时长", "备注"]),
    )
    write_main_sheet(
        wb,
        LINE_HEAD_SHEET,
        "EA118列头柜超功率告警统计表",
        ["序号", "数据时间", "机房", "楼栋", "房间", "机列", "功率", "对侧机列", "对侧机列最大功率", "次数", "时长", "备注"],
        rows_to_matrix(
            line_head_rows,
            ["序号", "数据时间", "机房", "楼栋", "房间", "机列", "功率", "对侧机列", "对侧机列最大功率", "次数", "时长", "备注"],
        ),
    )
    write_main_sheet(
        wb,
        CABINET_SHEET,
        "EA118机柜超功率告警统计表",
        ["序号", "数据时间", "机房", "楼栋", "房间", "机柜号", "机柜功率", "PDU编号", "电流值", "是否负载不均匀", "次数", "时长", "备注"],
        cabinet_groups_to_matrix(cabinet_groups),
        merge_row_spans=[max(1, len(group.get("pdu_rows") or [])) for group in cabinet_groups],
    )
    write_main_sheet(
        wb,
        BRANCH_SHEET,
        "EA118单支路超功率告警统计表",
        ["序号", "数据时间", "机房", "楼栋", "房间", "支路号", "PDU编号", "支路功率", "对侧PDU编号", "支路功率", "时长", "备注"],
        rows_to_matrix(
            branch_rows,
            ["序号", "数据时间", "机房", "楼栋", "房间", "支路号", "PDU编号", "支路功率", "对侧PDU编号", "对侧支路功率", "时长", "备注"],
        ),
    )
    write_info_sheet(wb)
    write_appendix_sheet(wb, ROW_LINE_APPENDIX_SHEET, row_line_source_rows, ("楼栋", "房间", "机列"), "功率", "KW")
    write_appendix_sheet(wb, LINE_HEAD_APPENDIX_SHEET, line_head_source_rows, ("楼栋", "房间", "机列"), "功率", "kw")
    write_appendix_sheet(wb, CABINET_APPENDIX_SHEET, filtered.get("cabinet", []), ("楼栋", "房间", "机柜号"), "机柜功率", "kw")

    wb.save(out)
    wb.close()
    return {
        "output_path": out,
        "counts": {
            "row_line": len(row_line_rows),
            "line_head": len(line_head_rows),
            "cabinet": len(cabinet_groups),
            "branch": len(branch_rows),
        },
        "date_label": date_label,
    }


def aggregate_simple_rows(
    rows: Sequence[Mapping[str, Any]],
    key_fields: Sequence[str],
    power_field: str,
    output_fields: Sequence[str],
    date_label: str,
    data_center: str,
) -> list[dict[str, Any]]:
    groups = group_rows(rows, key_fields)
    output: list[dict[str, Any]] = []
    for group in sorted(groups.values(), key=sort_group):
        max_row = max(group, key=lambda row: number_value(row.get(power_field)))
        item = {field: max_row.get(field) for field in output_fields}
        item["序号"] = len(output) + 1
        item["数据时间"] = date_label
        item["机房"] = item.get("机房") or data_center
        item["楼栋"] = normalize_building(item.get("楼栋"))
        total_hours = sum_hours(row.get("时长") for row in group)
        item["次数"] = sum_count_with_duration(group, total_hours)
        item["时长"] = format_hours(total_hours)
        output.append(item)
    return output


def aggregate_branch_rows(
    rows: Sequence[Mapping[str, Any]],
    date_label: str,
    data_center: str,
) -> list[dict[str, Any]]:
    normalized_rows = []
    for row in rows:
        normalized = dict(row)
        normalized["支路号"] = row.get("支路编号") or row.get("支路号")
        normalized["对侧支路功率"] = row.get("对侧支路功率") or row.get("支路功率_2") or row.get("支路功率.1")
        normalized_rows.append(normalized)

    groups = group_rows(normalized_rows, ("楼栋", "房间", "支路号", "PDU编号"))
    output: list[dict[str, Any]] = []
    for group in sorted(groups.values(), key=sort_group):
        max_row = max(group, key=lambda row: number_value(row.get("支路功率")))
        output.append(
            {
                "序号": len(output) + 1,
                "数据时间": date_label,
                "机房": max_row.get("机房") or data_center,
                "楼栋": normalize_building(max_row.get("楼栋")),
                "房间": max_row.get("房间"),
                "支路号": max_row.get("支路号"),
                "PDU编号": max_row.get("PDU编号"),
                "支路功率": max_row.get("支路功率"),
                "对侧PDU编号": max_row.get("对侧PDU编号"),
                "对侧支路功率": max_row.get("对侧支路功率"),
                "时长": format_hours(sum_hours(row.get("时长") for row in group)),
                "备注": max_row.get("备注"),
            }
        )
    return output


def aggregate_cabinet_rows(
    rows: Sequence[Mapping[str, Any]],
    date_label: str,
    data_center: str,
) -> list[dict[str, Any]]:
    groups = group_rows(rows, ("楼栋", "房间", "机柜号"))
    output: list[dict[str, Any]] = []
    for group in sorted(groups.values(), key=sort_group):
        max_row = max(group, key=lambda row: number_value(row.get("机柜功率")))
        max_date = parse_date(max_row.get("数据时间"))
        pdu_rows = [
            row
            for row in group
            if parse_date(row.get("数据时间")) == max_date
        ]
        pdu_rows.sort(key=lambda row: natural_key(str(row.get("PDU编号") or "")))
        if not pdu_rows:
            pdu_rows = [max_row]

        date_representatives: dict[date, Mapping[str, Any]] = {}
        for row in group:
            parsed = parse_date(row.get("数据时间"))
            if not parsed:
                continue
            current = date_representatives.get(parsed)
            if current is None or number_value(row.get("机柜功率")) > number_value(current.get("机柜功率")):
                date_representatives[parsed] = row

        representative_rows = list(date_representatives.values())
        total_hours = sum_hours(row.get("时长") for row in representative_rows)
        output.append(
            {
                "序号": len(output) + 1,
                "数据时间": date_label,
                "机房": max_row.get("机房") or data_center,
                "楼栋": normalize_building(max_row.get("楼栋")),
                "房间": max_row.get("房间"),
                "机柜号": max_row.get("机柜号"),
                "机柜功率": max_row.get("机柜功率"),
                "是否负载不均匀": max_row.get("是否负载不均匀") or "均匀",
                "次数": sum_count_with_duration(representative_rows, total_hours),
                "时长": format_hours(total_hours),
                "备注": max_row.get("备注"),
                "pdu_rows": pdu_rows,
            }
        )
    return output


def rows_to_matrix(rows: Sequence[Mapping[str, Any]], fields: Sequence[str]) -> list[list[Any]]:
    return [[row.get(field) for field in fields] for row in rows]


def cabinet_groups_to_matrix(groups: Sequence[Mapping[str, Any]]) -> list[list[Any]]:
    matrix: list[list[Any]] = []
    for group in groups:
        pdu_rows = list(group.get("pdu_rows") or [])
        for index, pdu_row in enumerate(pdu_rows):
            if index == 0:
                matrix.append(
                    [
                        group.get("序号"),
                        group.get("数据时间"),
                        group.get("机房"),
                        group.get("楼栋"),
                        group.get("房间"),
                        group.get("机柜号"),
                        group.get("机柜功率"),
                        pdu_row.get("PDU编号"),
                        pdu_row.get("电流值"),
                        group.get("是否负载不均匀"),
                        group.get("次数"),
                        group.get("时长"),
                        group.get("备注"),
                    ]
                )
            else:
                matrix.append([None, None, None, None, None, None, None, pdu_row.get("PDU编号"), pdu_row.get("电流值"), None, None, None, None])
    return matrix


def write_main_sheet(
    wb: Workbook,
    sheet_name: str,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    merge_row_spans: Sequence[int] | None = None,
) -> None:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        write_template_main_sheet(ws, title, headers, rows, merge_row_spans)
        return

    ws = wb.create_sheet(sheet_name)
    ws.append([title, *([None] * (len(headers) - 1))])
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    style_report_sheet(ws, len(headers))
    if merge_row_spans:
        merge_group_rows(ws, 3, merge_row_spans)


def write_template_main_sheet(
    ws: Any,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    merge_row_spans: Sequence[int] | None = None,
) -> None:
    column_count = len(headers)
    template_rows = capture_template_rows(ws, column_count, merge_row_spans)
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    if ws.max_row > 2:
        ws.delete_rows(3, ws.max_row - 2)
    trim_extra_columns(ws, column_count)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)

    ws.cell(1, 1).value = title
    for column_index, header in enumerate(headers, start=1):
        ws.cell(2, column_index).value = header
    set_font_name_for_range(ws, 1, 2, column_count)

    for row_offset, values in enumerate(rows):
        target_row = 3 + row_offset
        template = template_rows[row_offset % len(template_rows)]
        apply_template_row(ws, target_row, template, column_count)
        for column_index in range(1, column_count + 1):
            ws.cell(target_row, column_index).value = values[column_index - 1] if column_index <= len(values) else None

    if merge_row_spans:
        merge_group_rows(ws, 3, merge_row_spans)


def capture_template_rows(ws: Any, column_count: int, merge_row_spans: Sequence[int] | None = None) -> list[dict[str, Any]]:
    if merge_row_spans:
        start_row = first_visible_data_row(ws, column_count)
        row_count = max(1, min(max(merge_row_spans), max(ws.max_row - start_row + 1, 1)))
        return [capture_template_row(ws, min(start_row + index, ws.max_row), column_count) for index in range(row_count)]
    return [capture_template_row(ws, first_visible_data_row(ws, column_count), column_count)]


def first_visible_data_row(ws: Any, column_count: int) -> int:
    for row_index in range(3, ws.max_row + 1):
        if ws.row_dimensions[row_index].hidden:
            continue
        if any(ws.cell(row_index, column_index).value not in (None, "") for column_index in range(1, column_count + 1)):
            return row_index
    return 3


def capture_template_row(
    ws: Any,
    row_index: int,
    column_count: int,
    repeat_last_column_style: bool = False,
) -> dict[str, Any]:
    max_template_column = max(ws.max_column, 1)
    return {
        "height": ws.row_dimensions[row_index].height,
        "cells": [
            {
                "style": copy(ws.cell(row_index, min(column_index, max_template_column) if repeat_last_column_style else column_index)._style),
                "number_format": ws.cell(row_index, min(column_index, max_template_column) if repeat_last_column_style else column_index).number_format,
                "alignment": copy(ws.cell(row_index, min(column_index, max_template_column) if repeat_last_column_style else column_index).alignment),
            }
            for column_index in range(1, column_count + 1)
        ],
    }


def apply_template_row(ws: Any, row_index: int, template: Mapping[str, Any], column_count: int) -> None:
    ws.row_dimensions[row_index].height = template["height"]
    ws.row_dimensions[row_index].hidden = False
    for column_index in range(1, column_count + 1):
        cell = ws.cell(row_index, column_index)
        cell_template = template["cells"][column_index - 1]
        cell._style = copy(cell_template["style"])
        cell.number_format = cell_template["number_format"]
        cell.alignment = copy(cell_template["alignment"])
        set_cell_font_name(cell)


def write_template_matrix_sheet(
    ws: Any,
    values: Sequence[Sequence[Any]],
    header_row_count: int,
    preserved_merges: Sequence[tuple[int, int, int, int]] | None = None,
) -> None:
    column_count = max((len(row) for row in values), default=1)
    max_template_column = max(ws.max_column, 1)
    column_widths = [
        ws.column_dimensions[get_column_letter(min(column_index, max_template_column))].width
        for column_index in range(1, column_count + 1)
    ]
    template_max_row = max(ws.max_row, 1)
    header_templates = [
        capture_template_row(ws, min(row_index, template_max_row), column_count, repeat_last_column_style=True)
        for row_index in range(1, max(header_row_count, 1) + 1)
    ]
    data_template = capture_template_row(
        ws,
        min(header_row_count + 1, template_max_row),
        column_count,
        repeat_last_column_style=True,
    )

    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)

    for row_offset, row_values in enumerate(values, start=1):
        template = header_templates[row_offset - 1] if row_offset <= len(header_templates) else data_template
        apply_template_row(ws, row_offset, template, column_count)
        for column_index in range(1, column_count + 1):
            ws.cell(row_offset, column_index).value = row_values[column_index - 1] if column_index <= len(row_values) else None

    for column_index, width in enumerate(column_widths, start=1):
        if width is not None:
            ws.column_dimensions[get_column_letter(column_index)].width = width

    max_row = len(values)
    if preserved_merges:
        for start_row, start_column, end_row, end_column in preserved_merges:
            if max_row >= end_row and column_count >= end_column:
                ws.merge_cells(
                    start_row=start_row,
                    start_column=start_column,
                    end_row=end_row,
                    end_column=end_column,
                )
    trim_extra_columns(ws, column_count)


def set_cell_font_name(cell: Any) -> None:
    if isinstance(cell, MergedCell):
        return
    font = copy(cell.font)
    font.name = DEFAULT_FONT_NAME
    cell.font = font


def set_font_name_for_range(ws: Any, min_row: int, max_row: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col):
        for cell in row:
            set_cell_font_name(cell)


def trim_extra_columns(ws: Any, column_count: int) -> None:
    extra_column_count = ws.max_column - column_count
    if extra_column_count > 0:
        ws.delete_cols(column_count + 1, extra_column_count)


def merge_group_rows(ws: Any, start_row: int, merge_row_spans: Sequence[int]) -> None:
    row_index = start_row
    merge_columns = list(range(1, 8)) + list(range(10, 14))
    for span in merge_row_spans:
        span = max(1, int(span or 1))
        if span > 1:
            for column_index in merge_columns:
                ws.merge_cells(
                    start_row=row_index,
                    start_column=column_index,
                    end_row=row_index + span - 1,
                    end_column=column_index,
                )
        row_index += span


def write_info_sheet(wb: Workbook) -> None:
    if INFO_SHEET in wb.sheetnames:
        ws = wb[INFO_SHEET]
        for merged_range in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merged_range))
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(INFO_SHEET)
    rows = [
        ["填写说明", None],
        ["隐藏部分", "为历史超功率告警（本月未触发超功率），勿删除！"],
        [None, None],
        ["生成说明", "本表由四类超功率多维表日记录聚合生成。"],
        ["阈值说明", "机列：A/B/C/D楼 >215KW，E楼 >235KW；列头柜：A/B/C/D楼 >107.5KW，E楼 >117.5KW。"],
    ]
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 62
    ws["A1"].font = Font(bold=True)
    set_font_name_for_range(ws, 1, ws.max_row, ws.max_column)


def write_appendix_sheet(
    wb: Workbook,
    sheet_name: str,
    rows: Sequence[Mapping[str, Any]],
    key_fields: Sequence[str],
    power_field: str,
    unit: str,
) -> None:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        has_template = True
    else:
        ws = wb.create_sheet(sheet_name)
        has_template = False
    groups = group_rows(rows, key_fields)
    sorted_groups = sorted(groups.values(), key=sort_group)
    labels = [appendix_label(group[0], key_fields) for group in sorted_groups]
    output_rows: list[list[Any]] = [
        ["序号", "时间", *labels],
        [None, None, *("日最大功率" for _ in labels)],
        [None, None, *(unit for _ in labels)],
    ]

    dates = sorted({parsed for row in rows if (parsed := parse_date(row.get("数据时间"))) is not None})
    for index, current_date in enumerate(dates, start=1):
        values = []
        for group in sorted_groups:
            day_rows = [row for row in group if parse_date(row.get("数据时间")) == current_date]
            if not day_rows:
                values.append(None)
            else:
                values.append(max(number_value(row.get(power_field)) for row in day_rows))
        output_rows.append([index, format_day(current_date), *values])

    if has_template:
        write_template_matrix_sheet(
            ws,
            output_rows,
            header_row_count=3,
            preserved_merges=[(1, 1, 3, 1), (1, 2, 3, 2)],
        )
    else:
        for row in output_rows:
            ws.append(row)
        style_appendix_sheet(ws)


def style_report_sheet(ws: Any, column_count: int) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="E2F0D9")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    ws.freeze_panes = "A3"

    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 2), max_col=column_count):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.fill = title_fill
                cell.font = Font(bold=True, size=14)
            elif cell.row == 2:
                cell.fill = header_fill
                cell.font = Font(bold=True)
            set_cell_font_name(cell)

    for column_index in range(1, column_count + 1):
        letter = get_column_letter(column_index)
        width = 12
        if column_index in {2, 5, 6, 8, 9}:
            width = 18
        if column_index in {10, 11, 12, 13}:
            width = 14
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 24


def style_appendix_sheet(ws: Any) -> None:
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.row <= 3:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
            set_cell_font_name(cell)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 14
    for column_index in range(3, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(column_index)].width = 24


def group_rows(rows: Sequence[Mapping[str, Any]], key_fields: Sequence[str]) -> dict[tuple[str, ...], list[Mapping[str, Any]]]:
    groups: dict[tuple[str, ...], list[Mapping[str, Any]]] = defaultdict(list)
    for row in rows:
        key = tuple(str(row.get(field) or "").strip() for field in key_fields)
        if all(key):
            groups[key].append(row)
    return groups


def sort_group(group: Sequence[Mapping[str, Any]]) -> tuple[Any, ...]:
    first = group[0]
    return natural_key(" ".join(str(first.get(field) or "") for field in ("楼栋", "房间", "机列", "机柜号", "支路号", "PDU编号")))


def appendix_label(row: Mapping[str, Any], key_fields: Sequence[str]) -> str:
    parts = []
    for field in key_fields:
        value = row.get(field)
        if field == "楼栋":
            value = normalize_building(value)
        text = str(value or "").strip()
        if text:
            parts.append(text)
    return "/".join(parts)


def normalize_building(value: Any) -> str:
    text = str(value or "").strip()
    return text[:-1] if text.endswith("楼") else text


def filter_building_power_threshold(
    rows: Sequence[Mapping[str, Any]],
    power_field: str,
    default_threshold: float,
    e_building_threshold: float,
) -> list[dict[str, Any]]:
    filtered_rows: list[dict[str, Any]] = []
    for row in rows:
        building = normalize_building(row.get("楼栋")).upper()
        threshold = e_building_threshold if building == "E" else default_threshold
        power = number_value(row.get(power_field))
        if power != float("-inf") and power > threshold:
            filtered_rows.append(dict(row))
    return filtered_rows


def month_bounds(year: str, month_number: int) -> tuple[date, date]:
    start = date(int(year), int(month_number), 1)
    last = calendar.monthrange(start.year, start.month)[1]
    return start, date(start.year, start.month, last)


def date_in_range(value: date | None, start: date, end: date) -> bool:
    return bool(value and start <= value <= end)


def parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        raw = float(value)
        if raw > 10_000_000_000:
            return datetime.fromtimestamp(raw / 1000).date()
        if raw > 10_000_000:
            return datetime.fromtimestamp(raw).date()
    text = str(value).strip()
    if not text:
        return None
    match = re.search(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", text)
    if not match:
        return None
    return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))


def format_day(value: date) -> str:
    return f"{value.year}/{value.month}/{value.day}"


def number_value(value: Any) -> float:
    if value in (None, ""):
        return float("-inf")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).replace(",", "")
    match = HOUR_RE.search(text)
    return float(match.group(0)) if match else float("-inf")


def sum_int(values: Sequence[Any]) -> int:
    total = 0
    for value in values:
        number = number_value(value)
        if number != float("-inf"):
            total += int(number)
    return total


def sum_count_with_duration(rows: Sequence[Mapping[str, Any]], duration_hours: float | None = None) -> int:
    total = sum_int(row.get("次数") for row in rows)
    if total > 0:
        return total

    hours = duration_hours
    if hours is None:
        hours = sum_hours(row.get("时长") for row in rows)
    if rows and hours > 0:
        return 1
    return 0


def sum_hours(values: Sequence[Any]) -> float:
    total = 0.0
    for value in values:
        number = number_value(value)
        if number != float("-inf"):
            total += number
    return total


def format_hours(value: float) -> str:
    if float(value).is_integer():
        return f"{int(value)}h"
    return f"{value:.2f}".rstrip("0").rstrip(".") + "h"


def natural_key(value: str) -> tuple[Any, ...]:
    return tuple(int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value))
