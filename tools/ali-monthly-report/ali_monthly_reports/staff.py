from __future__ import annotations

import re
from itertools import zip_longest
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .utils import cell_text, date_text, first_present, is_truthy_cell, is_value_in_month, month_bounds


def is_formal_staff_record(record: dict[str, Any]) -> bool:
    raw_value = first_present(record, "岗位性质")
    values = [
        item.strip()
        for item in re.split(r"[\n,，;；/、]+", cell_text(raw_value))
        if item.strip()
    ]
    return "正式岗" in values


def format_staff_sheet(ws, widths: list[int]) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def is_intern_or_backup_position(position: str) -> bool:
    return "实习" in position or "备岗" in position


def write_staff_roster_workbook(
    active_records: list[dict[str, Any]],
    all_records: list[dict[str, Any]],
    month: str,
    output_path: Path,
) -> dict[str, int]:
    month_start, month_end = month_bounds(month)
    formal_active_records = [record for record in active_records if is_formal_staff_record(record)]
    formal_all_records = [record for record in all_records if is_formal_staff_record(record)]
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    active_ws = wb.active
    active_ws.title = "在职人员基础信息"
    active_ws.append(["人员姓名", "岗位", "入职时间"])
    for record in formal_active_records:
        active_ws.append(
            [
                cell_text(first_present(record, "姓名", "人员姓名", "员工姓名")),
                cell_text(first_present(record, "岗位")),
                date_text(first_present(record, "入职日期", "入职时间")),
            ]
        )

    new_hires: list[tuple[str, str]] = []
    leavers: list[tuple[str, str]] = []
    for record in all_records:
        name = cell_text(first_present(record, "姓名", "人员姓名", "员工姓名"))
        if not name:
            continue
        position = cell_text(first_present(record, "岗位"))

        if is_value_in_month(first_present(record, "入职日期", "入职时间"), month_start, month_end):
            new_hires.append((name, position))
        if is_truthy_cell(first_present(record, "离职/异动情况", "离职情况")) and is_value_in_month(
            first_present(record, "离职/异动日期", "离职日期"),
            month_start,
            month_end,
        ):
            leavers.append((name, position))

    intern_or_backup_new_hires = [entry for entry in new_hires if is_intern_or_backup_position(entry[1])]
    intern_or_backup_leavers = [entry for entry in leavers if is_intern_or_backup_position(entry[1])]
    other_new_hires = [entry for entry in new_hires if not is_intern_or_backup_position(entry[1])]
    other_leavers = [entry for entry in leavers if not is_intern_or_backup_position(entry[1])]

    changes_ws = wb.create_sheet("本月人员变动")
    title_rows: set[int] = set()
    header_rows: set[int] = set()
    for section_index, (title, section_new_hires, section_leavers) in enumerate(
        [
            ("实习生及备岗", intern_or_backup_new_hires, intern_or_backup_leavers),
            ("示例人员004员", other_new_hires, other_leavers),
        ]
    ):
        changes_ws.append([title])
        title_row = changes_ws.max_row
        title_rows.add(title_row)
        changes_ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=4)
        title_cell = changes_ws.cell(title_row, 1)
        title_cell.font = Font(bold=True, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor="1F4E78")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        changes_ws.append(["本月新增人员姓名", "岗位", "离职人员姓名", "岗位"])
        header_row = changes_ws.max_row
        header_rows.add(header_row)
        for cell in changes_ws[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for new_hire, leaver in zip_longest(section_new_hires, section_leavers, fillvalue=("", "")):
            changes_ws.append([new_hire[0], new_hire[1], leaver[0], leaver[1]])

        if section_index == 0:
            changes_ws.append([])

    for row in changes_ws.iter_rows():
        for cell in row:
            if cell.row not in title_rows | header_rows:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    changes_ws.freeze_panes = "A2"
    for index, width in enumerate([22, 24, 22, 24], start=1):
        changes_ws.column_dimensions[get_column_letter(index)].width = width

    format_staff_sheet(active_ws, [18, 24, 14])
    wb.save(output_path)
    wb.close()
    return {
        "active_count": len(formal_active_records),
        "new_hire_count": len(new_hires),
        "leaver_count": len(leavers),
        "ignored_non_formal_count": max(len(all_records) - len(formal_all_records), 0),
    }
